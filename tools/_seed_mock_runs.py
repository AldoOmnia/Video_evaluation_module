"""Seed the reshim dashboard with plausible mock runs.

Posts through the real ingest endpoint, so this exercises the same path the
daily workflow uses rather than writing files behind the server's back.

  .venv/bin/python tools/_seed_mock_runs.py http://localhost:3011 <token> [days]

Every figure is invented. Point it at a demo instance, not at a host whose
dashboard someone reads as production truth.
"""
import base64
import json
import random
import sys
import urllib.error
import urllib.request
from datetime import date, timedelta
from io import BytesIO

if len(sys.argv) < 3:
    sys.exit(__doc__)
base, token = sys.argv[1].rstrip("/"), sys.argv[2]
days = int(sys.argv[3]) if len(sys.argv) > 3 else 30

random.seed(425)  # stable output, so re-seeding does not reshuffle the history

FAMILIES = ["425", "430", "440", "450"]
VARIANTS = {
    "425": ["425-A-RH", "425-A-LH", "425-B-RH"],
    "430": ["430-B-LH", "430-C-RH"],
    "440": ["440-C-RH", "440-D-LH"],
    "450": ["450-E-RH"],
}

# Per-variant baseline BAD rate. Uniform randomness makes every family look
# alike and leaves the "variants flagged" panel permanently empty, which
# demonstrates nothing — the point of the report is that trouble concentrates.
# 440-D-LH is the persistent offender; 430-C-RH drifts and recovers.
BAD_RATE = {
    "425-A-RH": 0.07, "425-A-LH": 0.08, "425-B-RH": 0.11,
    "430-B-LH": 0.09, "430-C-RH": 0.14,
    "440-C-RH": 0.12, "440-D-LH": 0.34,
    "450-E-RH": 0.06,
}
# Share of failures that are BAD_HEAVY rather than BAD.
HEAVY_SHARE = 0.28


def workbook(day: date, rows: list[tuple]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "MOCK — sample data"
    ws["A1"] = f"MOCK DATA FOR {day.isoformat()} — NOT REAL MEASUREMENTS"
    ws["A1"].font = Font(bold=True, size=13, color="C00000")
    ws.merge_cells("A1:F1")
    ws.append([])
    ws.append(["Serial number", "Family", "Variant", "Backlash (mm)", "Spec", "Result"])
    for c in ws[3]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
    for col, w in zip("ABCDEF", (18, 10, 14, 15, 12, 12)):
        ws.column_dimensions[col].width = w
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_day(day: date, drift: float) -> dict:
    """One day of counts. `drift` lifts every variant's BAD rate together, so the
    sparkline shows a deterioration that recovers rather than pure noise."""
    by_family: dict[str, dict[str, int]] = {}
    by_variant: dict[str, dict[str, int]] = {}
    rows: list[tuple] = []
    ok = bad = heavy = 0

    for fam in FAMILIES:
        # ~120 units/day across the four families, which is the order of
        # magnitude the line actually runs, and enough that each variant clears
        # the n>=8 minimum the flagging rule needs.
        n = random.randint(22, 38)
        for _ in range(n):
            var = random.choice(VARIANTS[fam])
            rate = min(0.9, BAD_RATE[var] + drift)
            if random.random() >= rate:
                res, meas = "OK", round(random.uniform(0.14, 0.22), 3)
                ok += 1
            elif random.random() < HEAVY_SHARE:
                res, meas = "BAD_HEAVY", round(random.uniform(0.33, 0.45), 3)
                heavy += 1
            else:
                res, meas = "BAD", round(random.uniform(0.24, 0.30), 3)
                bad += 1
            by_family.setdefault(fam, {}).setdefault(res, 0)
            by_family[fam][res] += 1
            by_variant.setdefault(var, {}).setdefault(res, 0)
            by_variant[var][res] += 1
            if len(rows) < 40:
                rows.append((f"MOCK-{day:%m%d}-{len(rows):04d}", fam, var, meas, "0.13–0.23", res))

    high = []
    for var, counts in by_variant.items():
        n = sum(counts.values())
        b = counts.get("BAD", 0) + counts.get("BAD_HEAVY", 0)
        if n >= 8 and 100 * b / n >= 30:
            high.append({"variant": var, "n": n, "bad_pct": round(100 * b / n, 1)})

    total = ok + bad + heavy
    excluded = random.randint(2, 7)
    return {
        "date": day.isoformat(),
        # Tells the dashboard to show its "sample data" badge, so nobody reads
        # these numbers as the line's actual output.
        "mock": True,
        "summary": {
            "total": total + excluded,
            "excluded": excluded,
            "ok": ok,
            "bad": bad,
            "bad_heavy": heavy,
            "unknown_family": random.randint(0, 2),
            "by_family": by_family,
            "by_variant": by_variant,
            "high_bad_variants": sorted(high, key=lambda h: -h["bad_pct"])[:5],
        },
        "email": {
            "status": 202,
            "recipients": ["alfonso.guidone@walterscheid.com", "mattia_lugli@comerindustries.com"],
            "subject": f"Comer Fargo Reshim — {day.isoformat()} (MOCK)",
        },
        "report": {
            "name": f"MOCK-reshim-{day.isoformat()}.xlsx",
            "base64": base64.b64encode(workbook(day, rows)).decode(),
        },
    }


today = date.today()
sent = 0
for i in range(days - 1, -1, -1):
    day = today - timedelta(days=i)
    # A bulge of trouble around three weeks back that settles down again, so the
    # 30-day sparkline has a story in it instead of a flat band.
    drift = 0.16 * max(0.0, 1 - abs(i - 20) / 7.0)
    payload = build_day(day, drift)
    req = urllib.request.Request(
        f"{base}/api/reshim/runs",
        data=json.dumps(payload).encode(),
        method="POST",
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {token}"},
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            s = payload["summary"]
            incl = s["ok"] + s["bad"] + s["bad_heavy"]
            print(f"  {day} -> {r.status}  n={incl:<4} OK={100*s['ok']/incl:5.1f}%  flagged={len(s['high_bad_variants'])}")
            sent += 1
    except urllib.error.HTTPError as e:
        sys.exit(f"  {day} -> {e.code} {e.read().decode()[:200]}")

print(f"seeded {sent} runs into {base}")
