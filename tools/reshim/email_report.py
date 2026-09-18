"""Send daily reshim report via Microsoft Graph sendMail.

Reuses MSAL auth from sharepoint.py (Mail.Send scope). Plain-text body plus
xlsx attachment.
"""
from __future__ import annotations
import base64
from datetime import date
from pathlib import Path
from typing import Optional

import requests

from .analyze import ClassifiedSummary
from .config import Config
from .sharepoint import GRAPH_ROOT, _acquire_token


def _build_body(summary: ClassifiedSummary, day: date, sn_missing_photos: list[str]) -> str:
    included = summary.ok + summary.bad + summary.bad_heavy
    ok_pct = 100 * summary.ok / included if included else 0
    lines = [
        f"Comer Fargo — Reshim Daily Report",
        f"Analysis date: {day.isoformat()}",
        f"",
        f"— Summary —",
        f"  Total rows analyzed:  {summary.total}  (excluded: {summary.excluded})",
        f"  OK:                   {summary.ok}   ({ok_pct:.1f}%)",
        f"  BAD:                  {summary.bad}",
        f"  BAD_HEAVY:            {summary.bad_heavy}",
        f"  Unknown family:       {summary.unknown_family}",
        f"",
        f"— By Family —",
    ]
    for fam, cnt in sorted(summary.by_family.items()):
        total = sum(cnt.values())
        ok = cnt.get("OK", 0)
        pct = 100 * ok / total if total else 0
        lines.append(f"  Family {fam:<8}  n={total:<3}  OK={ok:<3}  BAD={cnt.get('BAD',0):<3}  BAD_HEAVY={cnt.get('BAD_HEAVY',0):<3}  ({pct:.1f}% OK)")
    lines.append("")
    if summary.high_bad_variants:
        lines.append("— Variants flagged (≥30% BAD) —")
        for v, n, p in summary.high_bad_variants:
            lines.append(f"  ⚠ {v:<24} n={n}  {p:.1f}% BAD")
        lines.append("")
    if sn_missing_photos:
        lines.append(f"— Missing photos ({len(sn_missing_photos)} SNs, review manually) —")
        for sn in sn_missing_photos[:20]:
            lines.append(f"  {sn}")
        if len(sn_missing_photos) > 20:
            lines.append(f"  … and {len(sn_missing_photos) - 20} more (see attachment)")
        lines.append("")
    lines.append("Full details in the attached workbook.")
    lines.append("")
    lines.append("— This report was generated automatically by the Daedalus reshim agent.")
    lines.append("   Reply to this email to reach Aldo.")
    return "\n".join(lines)


def _graph_send(cfg: Config, message: dict, *, interactive: bool) -> int:
    """POST a built message to Graph sendMail and return the status code."""
    # Mail.Send scope; note client-credentials also requires app-role granted in AAD
    token = _acquire_token(cfg, ["Mail.Send"], interactive=interactive)
    url = f"{GRAPH_ROOT}/users/{cfg.mail.from_addr}/sendMail"
    r = requests.post(url, json=message, headers={"Authorization": f"Bearer {token}"}, timeout=60)
    if r.status_code >= 300:
        raise RuntimeError(f"Graph sendMail failed: {r.status_code} {r.text}")
    return r.status_code


def send_report(
    cfg: Config,
    day: date,
    summary: ClassifiedSummary,
    xlsx_path: Path,
    sn_missing_photos: list[str],
    *,
    interactive: bool = False,
    override_recipients: Optional[list[str]] = None,
) -> dict:
    recipients = override_recipients or cfg.mail.recipients
    if not recipients:
        raise RuntimeError("No mail recipients configured (MAIL_RECIPIENTS)")

    subject = f"Comer Fargo Reshim — {day.isoformat()} ({summary.ok + summary.bad + summary.bad_heavy} units, {int(100*summary.ok/max(summary.ok+summary.bad+summary.bad_heavy,1))}% OK)"
    body = _build_body(summary, day, sn_missing_photos)

    attachment_bytes = xlsx_path.read_bytes()
    attachment_b64 = base64.standard_b64encode(attachment_bytes).decode("ascii")

    message = {
        "message": {
            "subject": subject,
            "body": {"contentType": "Text", "content": body},
            "toRecipients": [{"emailAddress": {"address": r}} for r in recipients],
            "replyTo": [{"emailAddress": {"address": cfg.mail.reply_to}}],
            "attachments": [{
                "@odata.type": "#microsoft.graph.fileAttachment",
                "name": xlsx_path.name,
                "contentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "contentBytes": attachment_b64,
            }],
        },
        "saveToSentItems": True,
    }

    status = _graph_send(cfg, message, interactive=interactive)
    return {"status": status, "recipients": recipients, "subject": subject, "size_kb": len(attachment_bytes) // 1024}


# ── Delivery self-test ────────────────────────────────────────────────────────

# Sample figures. Deliberately not round numbers, so the email looks like a
# real report to a reader skimming the layout, but every one is invented.
_TEST_FIGURES = {
    "total": 120, "excluded": 5, "ok": 98, "bad": 12, "bad_heavy": 5,
    "families": [("425", 61, 55), ("430", 34, 28), ("440", 20, 15)],
    "variant": ("425-A-RH", 14, 35.7),
}


def _build_test_body(day: date, sent_at: str) -> str:
    f = _TEST_FIGURES
    ok_pct = 100 * f["ok"] / (f["ok"] + f["bad"] + f["bad_heavy"])
    lines = [
        "*** TEST MESSAGE — SAMPLE DATA — NO ACTION REQUIRED ***",
        "",
        "This is a delivery test of the automated reshim report. It confirms the",
        "platform can compose and send the daily shim analysis to this distribution",
        "list. It was triggered manually.",
        "",
        "Every figure below, and every row in the attached workbook, is invented for",
        "this test. No shim stack was measured, no production data was read, and no",
        "quality issue is being reported. Please do not act on any of it. You can",
        "delete this email.",
        "",
        "— Sample figures (NOT REAL) —",
        f"  Total rows analyzed:  {f['total']}  (excluded: {f['excluded']})",
        f"  OK:                   {f['ok']}   ({ok_pct:.1f}%)",
        f"  BAD:                  {f['bad']}",
        f"  BAD_HEAVY:            {f['bad_heavy']}",
        "",
        "— Sample by family (NOT REAL) —",
    ]
    for fam, n, ok in f["families"]:
        lines.append(f"  Family {fam:<8}  n={n:<3}  OK={ok:<3}  ({100*ok/n:.1f}% OK)")
    v, vn, vp = f["variant"]
    lines += [
        "",
        "— Sample variant flag (NOT REAL) —",
        f"  ⚠ {v:<24} n={vn}  {vp:.1f}% BAD",
        "",
        "— What the real report contains —",
        "  • A daily OK / BAD / BAD_HEAVY breakdown of reshim events at station",
        "    130/135 (SHIMMING 1/2), classified against the Family-specific",
        "    backlash spec rather than a single global tolerance.",
        "  • Any variant running at or above 30% BAD, so a drifting family gets",
        "    attention before the rework cost lands.",
        "  • Operator photos matched to each serial number by OCR of the SN label,",
        "    and a list of serials whose photo is missing so they can be reviewed",
        "    by hand.",
        "  • The full row-level workbook attached, one row per unit, as here.",
        "",
        f"Nominal analysis date for this test: {day.isoformat()}",
        f"Sent: {sent_at}",
        "",
        "— Sent automatically by the Daedalus reshim agent.",
        "   Reply to this email to reach Aldo.",
    ]
    return "\n".join(lines)


def _build_test_workbook(day: date) -> tuple[str, bytes]:
    """A small workbook shaped like the real one, with obviously fake rows.

    Built here rather than through build_xlsx() because that needs a real
    analysis pass and operator photos; this only has to prove an attachment
    survives the trip.
    """
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "TEST — sample data"

    ws["A1"] = "TEST DATA — NOT REAL MEASUREMENTS — DO NOT USE"
    ws["A1"].font = Font(bold=True, size=13, color="C00000")
    ws.merge_cells("A1:F1")

    ws.append([])
    ws.append(["Serial number", "Family", "Variant", "Backlash (mm)", "Spec", "Result"])
    for c in ws[3]:
        c.font = Font(bold=True)

    for sn, fam, var, meas, spec, res in [
        ("TEST-0000001", "425", "425-A-RH", 0.18, "0.13–0.23", "OK"),
        ("TEST-0000002", "425", "425-A-RH", 0.27, "0.13–0.23", "BAD"),
        ("TEST-0000003", "430", "430-B-LH", 0.15, "0.12–0.22", "OK"),
        ("TEST-0000004", "440", "440-C-RH", 0.41, "0.14–0.24", "BAD_HEAVY"),
        ("TEST-0000005", "425", "425-A-RH", 0.19, "0.13–0.23", "OK"),
    ]:
        ws.append([sn, fam, var, meas, spec, res])

    ws.append([])
    ws.append(["Generated by the Daedalus reshim agent as a delivery test."])

    for col, w in zip("ABCDEF", (18, 10, 14, 15, 12, 12)):
        ws.column_dimensions[col].width = w

    buf = BytesIO()
    wb.save(buf)
    return f"TEST-reshim-{day.isoformat()}.xlsx", buf.getvalue()


def send_test_report(
    cfg: Config,
    day: date,
    *,
    interactive: bool = False,
    extra_recipients: Optional[list[str]] = None,
) -> dict:
    """Send a clearly-marked test report to verify mail delivery end to end.

    Exercises the same Graph app-only path and the same attachment handling the
    daily report uses, so a success here means the real report can also land.
    """
    from datetime import datetime, timezone

    recipients = list(cfg.mail.recipients)
    for r in extra_recipients or []:
        if r and r not in recipients:
            recipients.append(r)
    if not recipients:
        raise RuntimeError("No mail recipients configured (MAIL_RECIPIENTS)")

    sent_at = datetime.now(timezone.utc).astimezone().strftime("%Y-%m-%d %H:%M %Z")
    # The marker leads the subject so it is unmistakable in a notification
    # preview, where the body is never read.
    subject = f"[TEST — sample data] Comer Fargo Reshim — delivery test {day.isoformat()}"
    name, content = _build_test_workbook(day)

    message = {
        "message": {
            "subject": subject,
            "body": {"contentType": "Text", "content": _build_test_body(day, sent_at)},
            "toRecipients": [{"emailAddress": {"address": r}} for r in recipients],
            "replyTo": [{"emailAddress": {"address": cfg.mail.reply_to}}],
            "attachments": [{
                "@odata.type": "#microsoft.graph.fileAttachment",
                "name": name,
                "contentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "contentBytes": base64.standard_b64encode(content).decode("ascii"),
            }],
        },
        "saveToSentItems": True,
    }

    status = _graph_send(cfg, message, interactive=interactive)
    return {
        "status": status,
        "recipients": recipients,
        "subject": subject,
        "attachment": name,
        "size_kb": max(len(content) // 1024, 1),
    }
