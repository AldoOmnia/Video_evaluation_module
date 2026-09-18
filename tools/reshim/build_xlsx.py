"""v2-layout xlsx builder — classified rows + embedded images (twoCellAnchor)."""
from __future__ import annotations
import tempfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from PIL import Image
try:
    import pillow_heif
    pillow_heif.register_heif_opener()
except Exception:
    pass

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
from openpyxl.utils.units import pixels_to_EMU

from .analyze import ClassifiedRow, ClassifiedSummary
from .match import SnAssignment, MAX_IMGS_PER_SN


IMAGE_MAX_PX = 200
CELL_W_PX    = 210
CELL_H_PX    = 210
BATCH_TAG    = "b4"

HEADERS = [
    "Serial_Number", "Internal_Part_Number", "Family",
    "Pinion_mm", "Brake_mm", "Differential_mm",
    "Backlash_avg_mm", "Spec_Low", "Spec_High",
    "Backlash_Status", "Bucket",
    "Reshim_Date", "Image_Count", "Image_Source", "Match_Method",
    "Tag_Image", "Process_Image_1", "Process_Image_2",
    "Process_Image_3", "Process_Image_4", "Process_Image_5",
]


def _heic_to_jpg(src: Path, dst: Path, max_px: int = IMAGE_MAX_PX) -> None:
    img = Image.open(src)
    img.thumbnail((max_px, max_px), Image.LANCZOS)
    img.convert("RGB").save(dst, "JPEG", quality=70, optimize=True)


def build(
    rows: list[ClassifiedRow],
    assignments: dict[str, SnAssignment],
    summary: ClassifiedSummary,
    out_path: Path,
    unmatched_photos: int = 0,
    unknown_sn_photos: int = 0,
) -> Path:
    tmp = Path(tempfile.mkdtemp(prefix="reshim_jpg_"))
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Reshim + Images"
    ws1.append(HEADERS)

    img_col_start = HEADERS.index("Tag_Image") + 1
    img_col_end = img_col_start + MAX_IMGS_PER_SN - 1
    col_width_units = CELL_W_PX / 7.0
    for col in range(img_col_start, img_col_end + 1):
        ws1.column_dimensions[get_column_letter(col)].width = col_width_units

    rows_sorted = sorted(rows, key=lambda r: r.sn)
    for i, r in enumerate(rows_sorted):
        assn = assignments.get(r.sn)
        imgs = assn.photos if assn else []
        method = assn.match_method if assn else ""

        ws1.append([
            r.sn, r.internal_part_number, r.family,
            r.pinion_mm, r.brake_mm, r.differential_mm,
            r.backlash_avg_mm, r.spec_low, r.spec_high,
            r.backlash_status, r.bucket,
            r.reshim_date.isoformat() if r.reshim_date else "",
            len(imgs), BATCH_TAG if imgs else "", method,
            "", "", "", "", "", "",
        ])

        row_num = i + 2
        ws1.row_dimensions[row_num].height = CELL_H_PX * 0.75

        pad = 5
        for j, src in enumerate(imgs[:MAX_IMGS_PER_SN]):
            jpg = tmp / f"{i:04d}_{j}.jpg"
            _heic_to_jpg(src, jpg)
            xl_img = XLImage(str(jpg))
            col_idx = img_col_start + j
            xl_img.anchor = TwoCellAnchor(
                editAs="oneCell",
                _from=AnchorMarker(
                    col=col_idx - 1, colOff=pixels_to_EMU(pad),
                    row=row_num - 1, rowOff=pixels_to_EMU(pad),
                ),
                to=AnchorMarker(
                    col=col_idx, colOff=-pixels_to_EMU(pad),
                    row=row_num, rowOff=-pixels_to_EMU(pad),
                ),
            )
            ws1.add_image(xl_img)

    # Summary + Flags sheet
    ws2 = wb.create_sheet("Summary + Flags")
    w = ws2.append
    w(["Fargo Reshim Analysis — daily run"])
    w([])
    w(["Total rows (sorted ascending by SN)", summary.total])
    w(["Excluded", summary.excluded])
    w(["OK", summary.ok])
    w(["BAD", summary.bad])
    w(["BAD_HEAVY", summary.bad_heavy])
    w(["UNKNOWN_FAMILY", summary.unknown_family])
    w(["Photos matched (any method)", sum(1 for a in assignments.values() if a.photos)])
    w(["Photos with unmatched SN", unmatched_photos])
    w(["Photos of SNs not in DB pull", unknown_sn_photos])
    w([])
    w(["Family", "n", "OK", "BAD", "BAD_HEAVY", "%OK"])
    for fam, cnt in sorted(summary.by_family.items()):
        total = sum(cnt.values())
        ok = cnt.get("OK", 0)
        w([fam, total, ok, cnt.get("BAD", 0), cnt.get("BAD_HEAVY", 0), f"{100*ok/total:.1f}%" if total else "-"])
    w([])
    w(["Internal_Part_Number", "n", "OK", "BAD", "BAD_HEAVY", "%BAD"])
    for ipn, cnt in sorted(summary.by_variant.items()):
        total = sum(cnt.values())
        bad_pct = 100 * (cnt.get("BAD", 0) + cnt.get("BAD_HEAVY", 0)) / total if total else 0
        w([ipn, total, cnt.get("OK", 0), cnt.get("BAD", 0), cnt.get("BAD_HEAVY", 0), f"{bad_pct:.1f}%"])
    w([])
    w(["High-BAD variants (≥30% BAD, sorted worst first):"])
    for v, n, p in summary.high_bad_variants:
        w([v, n, f"{p:.1f}%"])
    w([])
    w(["SNs missing photos:"])
    for r in rows_sorted:
        if r.sn not in assignments or not assignments[r.sn].photos:
            w([r.sn, r.reshim_date.isoformat() if r.reshim_date else ""])

    wb.save(out_path)
    return out_path
